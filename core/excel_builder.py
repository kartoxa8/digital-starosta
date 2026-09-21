from io import BytesIO
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from core.models import Attendance, Group, Lesson, Student, Subject


def build_report(
    group: Group,
    students: list[Student],
    lessons: list[Lesson],
    subjects: dict[int, Subject],
    records: list[Attendance],
) -> BytesIO:
    book = Workbook()

    # Sheet 1: Master Attendance Matrix
    sheet = book.active
    sheet.title = "Attendance Matrix"

    # Title & Metadata
    sheet["A1"] = f"{group.name} — {group.university}"
    sheet["A1"].font = Font(name="Calibri", size=15, bold=True, color="1F497D")
    sheet["A2"] = "Master Group Attendance Register"
    sheet["A2"].font = Font(name="Calibri", size=11, italic=True, color="595959")

    # Header Row
    headers = ["Subgroup", "Student Full Name", "Classes Held", "Absences", "Attendance Rate"]
    for lesson in lessons:
        subj_name = subjects[lesson.subject_id].name if lesson.subject_id in subjects else "Unknown"
        sg_tag = f"Subgroup {lesson.subgroup}" if lesson.subgroup else "All"
        date_str = lesson.opened_at.strftime("%d.%m") if lesson.opened_at else "--.--"
        headers.append(f"{date_str} {subj_name} ({sg_tag})")

    sheet.append([])  # Row 3 is empty spacer
    sheet.append(headers)  # Row 4 is table header
    header_row_idx = 4

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    by_lesson_student: dict[tuple[int, int], Attendance] = {
        (rec.lesson_id, rec.student_id): rec for rec in records
    }

    current_row = header_row_idx + 1

    # Sort students by subgroup, then alphabetical by last name
    sorted_students = sorted(students, key=lambda s: (s.subgroup, s.last_name, s.first_name))

    for student in sorted_students:
        # Lessons applicable to this student (subgroup 0 = common, or student's own subgroup)
        applicable_lessons = [l for l in lessons if l.subgroup in (0, student.subgroup)]
        classes_held = len(applicable_lessons)

        present_count = 0
        lesson_marks = []
        for l in lessons:
            if l.subgroup not in (0, student.subgroup):
                lesson_marks.append("-")
            else:
                rec = by_lesson_student.get((l.id, student.id))
                if rec:
                    if rec.status == "present":
                        present_count += 1
                        lesson_marks.append("+")
                    elif rec.status == "late":
                        present_count += 1
                        lesson_marks.append("О")
                    elif rec.status == "excused":
                        lesson_marks.append("У")
                    else:
                        lesson_marks.append("Н")
                else:
                    lesson_marks.append("Н")

        absences = classes_held - present_count
        rate = (present_count / classes_held) if classes_held > 0 else 1.0

        full_name = f"{student.last_name} {student.first_name}" + (f" {student.middle_name}" if student.middle_name else "")
        row_data = [
            f"Subgroup {student.subgroup}",
            full_name,
            classes_held,
            absences,
            rate,
        ] + lesson_marks

        sheet.append(row_data)

        # Style data row
        for col_idx in range(1, len(row_data) + 1):
            cell = sheet.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            if col_idx == 5:
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (1, 3, 4) or col_idx > 5:
                cell.alignment = Alignment(horizontal="center")

        current_row += 1

    # Freeze panes and column dimensions
    sheet.freeze_panes = "C5"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 18

    for col in range(6, len(headers) + 1):
        col_letter = get_column_letter(col)
        sheet.column_dimensions[col_letter].width = 20

    # Conditional formatting on Attendance Rate (Column E)
    if current_row > header_row_idx + 1:
        data_end_row = current_row - 1
        rate_range = f"E5:E{data_end_row}"

        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        green_font = Font(color="006100", bold=True)
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        yellow_font = Font(color="9C6500", bold=True)
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_font = Font(color="9C0006", bold=True)

        sheet.conditional_formatting.add(
            rate_range,
            CellIsRule(operator="greaterThanOrEqual", formula=["0.80"], fill=green_fill, font=green_font),
        )
        sheet.conditional_formatting.add(
            rate_range,
            CellIsRule(operator="between", formula=["0.60", "0.799"], fill=yellow_fill, font=yellow_font),
        )
        sheet.conditional_formatting.add(
            rate_range,
            CellIsRule(operator="lessThan", formula=["0.60"], fill=red_fill, font=red_font),
        )

    # Sheet 2: Security & Audit Trail
    audit = book.create_sheet("Security & Audit Trail")
    audit["A1"] = f"Security Verification & GPS Audit Trail — {group.name}"
    audit["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F497D")

    audit_headers = [
        "Record ID",
        "Student Name",
        "Telegram ID",
        "Lesson ID",
        "Scanned At (UTC)",
        "Device Latitude",
        "Device Longitude",
        "Distance to Campus (m)",
        "Attendance Status",
    ]
    audit.append([])
    audit.append(audit_headers)
    audit_header_row = 3

    for col_idx in range(1, len(audit_headers) + 1):
        c = audit.cell(row=audit_header_row, column=col_idx)
        c.font = header_font
        c.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        c.alignment = Alignment(horizontal="center")

    students_map = {s.id: s for s in students}

    for item in sorted(records, key=lambda x: x.scanned_at or ""):
        student_obj = students_map.get(item.student_id)
        s_name = (
            f"{student_obj.last_name} {student_obj.first_name}"
            if student_obj
            else f"Student #{item.student_id}"
        )
        tg_id = student_obj.telegram_id if student_obj else None
        scanned_str = item.scanned_at.strftime("%Y-%m-%d %H:%M:%S") if item.scanned_at else ""

        audit.append([
            item.id,
            s_name,
            tg_id,
            item.lesson_id,
            scanned_str,
            item.latitude,
            item.longitude,
            round(item.distance_meters, 1) if item.distance_meters is not None else "N/A",
            item.status,
        ])

    audit.column_dimensions["A"].width = 12
    audit.column_dimensions["B"].width = 28
    audit.column_dimensions["C"].width = 16
    audit.column_dimensions["D"].width = 12
    audit.column_dimensions["E"].width = 22
    audit.column_dimensions["F"].width = 16
    audit.column_dimensions["G"].width = 16
    audit.column_dimensions["H"].width = 22
    audit.column_dimensions["I"].width = 18

    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output
